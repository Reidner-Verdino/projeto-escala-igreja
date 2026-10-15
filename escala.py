import calendar
import json
import random
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HOMENS = [
    "Rafael",
    "Carlos",
    "Pedro",
    "Lucas",
    "Daniel",
    "Marcos",
    "Eduardo",
    "André",
    "Felipe",
    "Gabriel",
]

TAREFAS = [
    "Escola Bíblica Dominical",
    "Oração dos Homens",
    "Louvores",
    "Avisos",
    "Slides",
    "Momento de Oração",
    "Culto",
    "Ceia",
]

TAREFAS_FIXAS = {
    "Escola Bíblica Dominical": "Pedro",
    "Culto": "Pedro",
}

RESTRICOES_PESSOA = {
    "Eduardo": {"Oração dos Homens"},
}

RESTRICOES_TAREFA = {
    "Louvores": {"Daniel", "Carlos", "André"},
    "Slides": {"Carlos", "Marcos", "Rafael"},
}


PESSOAS_PAO = ["Helena", "Felipe", "Lucas", "André"]

GRUPOS = [
    {"label": "André, Laura e Rafael", "membros": {"André", "Laura", "Rafael"}},
    {"label": "Lucas e Júlia", "membros": {"Lucas", "Júlia"}},
    {"label": "Marcos, Clara e Gabriel", "membros": {"Marcos", "Clara", "Gabriel"}},
    {"label": "Carlos, Sofia e Eduardo", "membros": {"Carlos", "Sofia", "Eduardo"}},
    {"label": "Felipe e Marina", "membros": {"Felipe", "Marina"}},
]

GRUPOS_EXCLUIDOS_DO_ALMOCO = {"Carlos, Sofia e Eduardo"}

PESO_TRIO = 2
PESO_DUPLA = 1

HOMENS_DE_FOLGA_POR_SEMANA = 1

ARQUIVO_ESTADO_HOMENS = Path("estado_homens.json")
ARQUIVO_ESTADO_IGREJA = Path("estado_igreja.json")

NOMES_MESES = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril",
    5: "maio", 6: "junho", 7: "julho", 8: "agosto",
    9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro",
}
MESES_POR_NOME = {nome: numero for numero, nome in NOMES_MESES.items()}


def carregar_estado(caminho: Path) -> dict:
    if caminho.exists():
        with open(caminho, "r", encoding="utf-8") as f:
            return json.load(f)
    return {"semanas": []}


def salvar_estado(estado: dict, caminho: Path) -> None:
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(estado, f, ensure_ascii=False, indent=2)


def remover_semanas_do_mes(semanas: list, mes: int, ano: int) -> list:
    restantes = []
    for semana in semanas:
        data = date.fromisoformat(semana["data"])
        if data.month == mes and data.year == ano:
            continue
        restantes.append(semana)
    return restantes


def elegiveis_para_tarefa(tarefa, pessoas, restricoes_pessoa, restricoes_tarefa):
    permitidos_pela_tarefa = restricoes_tarefa.get(tarefa)
    elegiveis = []
    for pessoa in pessoas:
        tarefas_da_pessoa = restricoes_pessoa.get(pessoa)
        pode_por_pessoa = tarefas_da_pessoa is None or tarefa in tarefas_da_pessoa
        pode_por_tarefa = permitidos_pela_tarefa is None or pessoa in permitidos_pela_tarefa
        if pode_por_pessoa and pode_por_tarefa:
            elegiveis.append(pessoa)
    return elegiveis


def embaralhar_com_peso(itens: list, pesos: dict) -> list:
    def chave(item):
        peso = pesos.get(item, 1)
        return random.random() ** (1.0 / peso)
    return sorted(itens, key=chave, reverse=True)


def quem_ja_fez(semanas: list) -> dict:
    resultado = {}
    for semana in semanas:
        for tarefa, valor in semana.items():
            pessoas = valor if isinstance(valor, list) else [valor]
            for pessoa in pessoas:
                resultado.setdefault(tarefa, set()).add(pessoa)
    return resultado


def _montar_semana(tarefas_ordenadas, pessoas, ja_fez, permitir_repeticao, ja_usados,
                    restricoes_pessoa, restricoes_tarefa, pesos):
    def backtrack(index, usados, escala):
        if index == len(tarefas_ordenadas):
            return dict(escala)

        tarefa = tarefas_ordenadas[index]
        candidatos = [
            p for p in elegiveis_para_tarefa(tarefa, pessoas, restricoes_pessoa, restricoes_tarefa)
            if p not in usados
            and (permitir_repeticao or p not in ja_fez.get(tarefa, set()))
        ]
        if pesos:
            candidatos = embaralhar_com_peso(candidatos, pesos)
        else:
            random.shuffle(candidatos)

        for pessoa in candidatos:
            escala[tarefa] = pessoa
            resultado = backtrack(index + 1, usados | {pessoa}, escala)
            if resultado is not None:
                return resultado
            del escala[tarefa]

        return None

    return backtrack(0, set(ja_usados), {})


def sortear_semana(pessoas, tarefas, ja_fez, restricoes_pessoa=None, restricoes_tarefa=None,
                    tarefas_fixas=None, pesos=None):
    restricoes_pessoa = restricoes_pessoa or {}
    restricoes_tarefa = restricoes_tarefa or {}
    tarefas_fixas = tarefas_fixas or {}
    pesos = pesos or {}

    tarefas_a_sortear = [t for t in tarefas if t not in tarefas_fixas]
    ja_usados = set(tarefas_fixas.values())

    tarefas_ordenadas = sorted(
        tarefas_a_sortear,
        key=lambda t: len(elegiveis_para_tarefa(t, pessoas, restricoes_pessoa, restricoes_tarefa)),
    )

    resultado = _montar_semana(
        tarefas_ordenadas, pessoas, ja_fez, False, ja_usados,
        restricoes_pessoa, restricoes_tarefa, pesos,
    )

    if resultado is None:
        print("Aviso: não foi possível evitar repetição nesta semana.")
        resultado = _montar_semana(
            tarefas_ordenadas, pessoas, ja_fez, True, ja_usados,
            restricoes_pessoa, restricoes_tarefa, pesos,
        )

    if resultado is None:
        raise ValueError("Não foi possível montar a escala da semana — confira as listas e restrições.")

    escala_completa = dict(tarefas_fixas)
    escala_completa.update(resultado)

    return {tarefa: escala_completa[tarefa] for tarefa in tarefas}


def homens_com_grupo() -> list:
    pessoas = set()
    for grupo in GRUPOS:
        pessoas |= (grupo["membros"] & set(HOMENS))
    pessoas -= set(TAREFAS_FIXAS.values())
    return sorted(pessoas)


def contagem_de_folgas(semanas_homens: list) -> dict:
    contagem = {}
    for semana in semanas_homens:
        for pessoa in semana.get("folga_designada", []):
            contagem[pessoa] = contagem.get(pessoa, 0) + 1
    return contagem


def escolher_folga_da_semana(elegiveis: list, quantidade: int, contagem_historica: dict) -> set:
    candidatos = elegiveis.copy()
    random.shuffle(candidatos)
    candidatos.sort(key=lambda p: contagem_historica.get(p, 0))
    return set(candidatos[:quantidade])


def domingos_do_mes(mes: int, ano: int) -> list:
    dias_no_mes = calendar.monthrange(ano, mes)[1]
    return [
        date(ano, mes, dia)
        for dia in range(1, dias_no_mes + 1)
        if date(ano, mes, dia).weekday() == 6
    ]


def _valor_para_exibir(valor):
    return ", ".join(valor) if isinstance(valor, list) else valor


def _escrever_homens_na_aba(ws, semanas_do_mes, mes, ano):
    nome_mes = NOMES_MESES[mes].capitalize()
    ws["A1"] = f"Escala dos Homens — {nome_mes} de {ano}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center")

    cabecalho_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    semana_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    linha = 3
    for i, semana in enumerate(semanas_do_mes, start=1):
        data_str = date.fromisoformat(semana["data"]).strftime("%d/%m")

        ws.cell(row=linha, column=1, value=f"Semana {i} — {data_str}").font = Font(bold=True)
        ws.cell(row=linha, column=1).fill = semana_fill
        ws.cell(row=linha, column=2).fill = semana_fill
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        linha += 1

        folga = semana.get("folga_designada") or []
        if folga:
            ws.cell(row=linha, column=1, value=f"Folga: {', '.join(folga)}").font = Font(italic=True, size=9)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
            linha += 1

        ws.cell(row=linha, column=1, value="Tarefa").font = Font(bold=True)
        ws.cell(row=linha, column=2, value="Responsável").font = Font(bold=True)
        ws.cell(row=linha, column=1).fill = cabecalho_fill
        ws.cell(row=linha, column=2).fill = cabecalho_fill
        linha += 1

        for tarefa, pessoa in semana["escala"].items():
            ws.cell(row=linha, column=1, value=tarefa)
            ws.cell(row=linha, column=2, value=_valor_para_exibir(pessoa))
            linha += 1

        linha += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26


def _escrever_igreja_na_aba(ws, semanas_do_mes, mes, ano):
    nome_mes = NOMES_MESES[mes].capitalize()
    ws["A1"] = f"Pão, Pré-Culto, Almoço e Pós-Culto — {nome_mes} de {ano}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center")

    cabecalho_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    semana_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    linha = 3
    for i, semana in enumerate(semanas_do_mes, start=1):
        data_str = date.fromisoformat(semana["data"]).strftime("%d/%m")

        ws.cell(row=linha, column=1, value=f"Semana {i} — {data_str}").font = Font(bold=True)
        ws.cell(row=linha, column=1).fill = semana_fill
        ws.cell(row=linha, column=2).fill = semana_fill
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        linha += 1

        ws.cell(row=linha, column=1, value="Tarefa").font = Font(bold=True)
        ws.cell(row=linha, column=2, value="Responsável").font = Font(bold=True)
        ws.cell(row=linha, column=1).fill = cabecalho_fill
        ws.cell(row=linha, column=2).fill = cabecalho_fill
        linha += 1

        for tarefa, chave in (("Pão", "pao"), ("Pré-Culto", "pre_culto"),
                              ("Almoço", "almoco"), ("Pós-Culto", "pos_culto")):
            ws.cell(row=linha, column=1, value=tarefa)
            ws.cell(row=linha, column=2, value=_valor_para_exibir(semana[chave]))
            linha += 1

        linha += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30


def exportar_mes_para_excel(semanas_homens, semanas_igreja, mes, ano) -> str:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Escala dos Homens"
    _escrever_homens_na_aba(ws1, semanas_homens, mes, ano)

    ws2 = wb.create_sheet("Pão Pré-Culto Almoço")
    _escrever_igreja_na_aba(ws2, semanas_igreja, mes, ano)

    nome_mes = NOMES_MESES[mes].capitalize()
    nome_arquivo = f"escala_{nome_mes.lower()}_{ano}.xlsx"
    wb.save(nome_arquivo)
    return nome_arquivo


def interpretar_mes(texto: str) -> int:
    texto_normalizado = texto.strip().lower()
    if texto_normalizado.isdigit():
        return int(texto_normalizado)
    if texto_normalizado in MESES_POR_NOME:
        return MESES_POR_NOME[texto_normalizado]
    raise ValueError(f"Não reconheci o mês '{texto}'.")


def mes_seguinte(hoje: date) -> tuple:
    if hoje.month == 12:
        return 1, hoje.year + 1
    return hoje.month + 1, hoje.year


def gerar_mes_inteiro(mes: int, ano: int) -> None:
    datas = domingos_do_mes(mes, ano)

    estado_homens = carregar_estado(ARQUIVO_ESTADO_HOMENS)
    estado_igreja = carregar_estado(ARQUIVO_ESTADO_IGREJA)

    estado_homens["semanas"] = remover_semanas_do_mes(estado_homens["semanas"], mes, ano)
    estado_igreja["semanas"] = remover_semanas_do_mes(estado_igreja["semanas"], mes, ano)

    candidatos_a_folga = homens_com_grupo()

    print(f"\nGerando escalas de {NOMES_MESES[mes]} de {ano} ({len(datas)} semana(s))\n" + "=" * 40)

    for i, data in enumerate(datas, start=1):
        print(f"\nSemana {i} — {data.strftime('%d/%m')}\n" + "-" * 40)

        historico_folga = contagem_de_folgas(estado_homens["semanas"])
        folga_designada = escolher_folga_da_semana(candidatos_a_folga, HOMENS_DE_FOLGA_POR_SEMANA, historico_folga)

        homens_disponiveis = [h for h in HOMENS if h not in folga_designada]
        ja_fez_homens = quem_ja_fez([s["escala"] for s in estado_homens["semanas"]])
        escala_homens = sortear_semana(
            homens_disponiveis, TAREFAS, ja_fez_homens,
            restricoes_pessoa=RESTRICOES_PESSOA,
            restricoes_tarefa=RESTRICOES_TAREFA,
            tarefas_fixas=TAREFAS_FIXAS,
        )

        ociosos = set(HOMENS) - set(escala_homens.values())

        semana_homens = {
            "data": data.isoformat(),
            "escala": escala_homens,
            "ociosos": sorted(ociosos),
            "folga_designada": sorted(folga_designada),
        }
        estado_homens["semanas"].append(semana_homens)

        print("Escala dos Homens:")
        for tarefa, pessoa in escala_homens.items():
            print(f"  {tarefa:<28} {pessoa}")
        print(f"  Folga designada: {', '.join(sorted(folga_designada)) or '(nenhuma)'}")
        print(f"  Ociosos (vão pro Pós-Culto): {', '.join(sorted(ociosos)) or '(nenhum)'}")

        grupos_pos_culto = sorted({
            grupo["label"] for grupo in GRUPOS if grupo["membros"] & ociosos
        })
        grupos_ativos = [
            grupo for grupo in GRUPOS if not (grupo["membros"] & ociosos)
        ]
        labels_ativos = [grupo["label"] for grupo in grupos_ativos]
        pesos_grupos = {
            grupo["label"]: (PESO_TRIO if len(grupo["membros"]) >= 3 else PESO_DUPLA)
            for grupo in grupos_ativos
        }
        labels_para_almoco = set(labels_ativos) - GRUPOS_EXCLUIDOS_DO_ALMOCO

        ja_fez_grupos = quem_ja_fez([
            {"Pré-Culto": s["pre_culto"], "Almoço": s["almoco"]}
            for s in estado_igreja["semanas"]
        ])
        escala_grupos = sortear_semana(
            labels_ativos, ["Pré-Culto", "Almoço"], ja_fez_grupos,
            restricoes_tarefa={"Almoço": labels_para_almoco},
            pesos=pesos_grupos,
        )

        pao_disponiveis = [p for p in PESSOAS_PAO if p not in ociosos]
        ja_fez_pao = quem_ja_fez([{"Pão": s["pao"]} for s in estado_igreja["semanas"]])
        escala_pao = sortear_semana(pao_disponiveis, ["Pão"], ja_fez_pao)

        semana_igreja = {
            "data": data.isoformat(),
            "pao": escala_pao["Pão"],
            "pre_culto": escala_grupos["Pré-Culto"],
            "almoco": escala_grupos["Almoço"],
            "pos_culto": grupos_pos_culto,
        }
        estado_igreja["semanas"].append(semana_igreja)

        print("Igreja:")
        print(f"  Pão                          {semana_igreja['pao']}")
        print(f"  Pré-Culto                    {semana_igreja['pre_culto']}")
        print(f"  Almoço                       {semana_igreja['almoco']}")
        print(f"  Pós-Culto                    {', '.join(semana_igreja['pos_culto']) or '(nenhum)'}")

    salvar_estado(estado_homens, ARQUIVO_ESTADO_HOMENS)
    salvar_estado(estado_igreja, ARQUIVO_ESTADO_IGREJA)

    semanas_homens_do_mes = [s for s in estado_homens["semanas"] if date.fromisoformat(s["data"]).month == mes and date.fromisoformat(s["data"]).year == ano]
    semanas_igreja_do_mes = [s for s in estado_igreja["semanas"] if date.fromisoformat(s["data"]).month == mes and date.fromisoformat(s["data"]).year == ano]

    arquivo = exportar_mes_para_excel(semanas_homens_do_mes, semanas_igreja_do_mes, mes, ano)
    print(f"\nPlanilha gerada: {arquivo}")


def main():
    if len(sys.argv) > 1:
        mes = interpretar_mes(sys.argv[1])
        ano = int(sys.argv[2]) if len(sys.argv) > 2 else date.today().year
    else:
        mes, ano = mes_seguinte(date.today())

    gerar_mes_inteiro(mes, ano)


if __name__ == "__main__":
    main()
